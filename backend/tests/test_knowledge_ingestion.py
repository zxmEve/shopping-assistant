from __future__ import annotations

import hashlib
import io
from datetime import timedelta
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image
from sqlalchemy import create_engine, func, select

from procure_pilot.api.app import create_app
from procure_pilot.config import Settings
from procure_pilot.domain.knowledge import KNOWLEDGE_EMBEDDING_DIMENSIONS
from procure_pilot.domain.models import utc_now
from procure_pilot.evals.runner import make_engine
from procure_pilot.persistence.knowledge_repository import SqlAlchemyKnowledgeRepository
from procure_pilot.persistence.models import Base, KnowledgeDocumentVersionRecord
from procure_pilot.services.document_ingestion import (
    DOCX_MEDIA_TYPE,
    DocumentParseError,
    PolicyFileParser,
)
from procure_pilot.services.embeddings import OpenAIEmbeddingProvider
from procure_pilot.services.knowledge import PolicyKnowledgeService

CATEGORY = "CAT-COMPUTER-LAPTOP"
ADMIN_SESSION = {"Authorization": "Bearer fixture-session-cn-knowledge-admin"}
EMPLOYEE_SESSION = {"Authorization": "Bearer fixture-session-cn-employee-001"}


class FixtureOcrProvider:
    name = "fixture-ocr"

    def extract_pdf_page(self, pdf_bytes: bytes, page_index: int) -> str:
        assert pdf_bytes.startswith(b"%PDF-")
        assert page_index == 0
        return "第一条 扫描件采购制度\n研发笔记本内存不得低于32GB，保修36个月。"


class ReplacementEmbeddingProvider:
    model = "fixture-semantic-v2"
    dimensions = KNOWLEDGE_EMBEDDING_DIMENSIONS

    def embed(self, texts):
        result = []
        for index, _ in enumerate(texts):
            vector = [0.0] * self.dimensions
            vector[index % self.dimensions] = 1.0
            result.append(vector)
        return result


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("研发设备采购制度", level=1)
    document.add_paragraph("第一条 研发笔记本不得低于32GB内存，保修不得少于36个月。")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "金额"
    table.rows[0].cells[1].text = "审批人"
    table.rows[1].cells[0].text = "超过10000元"
    table.rows[1].cells[1].text = "直属经理"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _scanned_pdf_bytes() -> bytes:
    image = Image.new("RGB", (800, 300), "white")
    output = io.BytesIO()
    image.save(output, format="PDF")
    return output.getvalue()


@pytest.fixture
def knowledge_repository(tmp_path):
    database_engine = create_engine(f"sqlite:///{tmp_path / 'knowledge-ingestion.db'}")
    Base.metadata.create_all(database_engine)
    return SqlAlchemyKnowledgeRepository("sqlite://", engine=database_engine)


def test_docx_parser_preserves_heading_and_table_structure():
    data = _docx_bytes()

    parsed = PolicyFileParser().parse(
        data,
        file_name="研发制度.docx",
        media_type=DOCX_MEDIA_TYPE,
    )

    assert "# 研发设备采购制度" in parsed.raw_content
    assert "## 表格 1" in parsed.raw_content
    assert "| 金额 | 审批人 |" in parsed.raw_content
    assert "| 超过10000元 | 直属经理 |" in parsed.raw_content
    assert parsed.table_count == 1
    assert parsed.file_hash == hashlib.sha256(data).hexdigest()


def test_scanned_pdf_requires_ocr_and_uses_injected_provider():
    data = _scanned_pdf_bytes()
    with pytest.raises(DocumentParseError) as missing:
        PolicyFileParser().parse(data, file_name="scan.pdf", media_type="application/pdf")
    assert missing.value.code == "OCR_REQUIRED"

    parsed = PolicyFileParser(ocr_provider=FixtureOcrProvider()).parse(
        data,
        file_name="scan.pdf",
        media_type="application/pdf",
    )
    assert parsed.page_count == 1
    assert parsed.ocr_pages == [1]
    assert parsed.warnings == ["ocr_provider:fixture-ocr"]
    assert "扫描件采购制度" in parsed.raw_content


def test_file_ingestion_persists_original_blob_and_extraction_metadata(knowledge_repository):
    service = PolicyKnowledgeService(knowledge_repository)
    data = _docx_bytes()

    result = service.ingest_file(
        tenant_id="demo_cn",
        actor_user_id="knowledge_admin_001",
        roles=("KNOWLEDGE_STEWARD",),
        document_key="docx-policy",
        title="研发设备采购制度",
        category_id=CATEGORY,
        source_uri="s3://policies/docx-policy.docx",
        policy_version="3",
        effective_from=utc_now() - timedelta(days=1),
        file_name="docx-policy.docx",
        media_type=DOCX_MEDIA_TYPE,
        file_data=data,
    )

    assert result.version.source_file_hash == hashlib.sha256(data).hexdigest()
    assert result.version.parser_version == "policy-file-parser-v1"
    assert result.version.extraction_metadata["table_count"] == 1
    with knowledge_repository.session_factory() as session:
        record = session.get(KnowledgeDocumentVersionRecord, result.version.version_id)
        assert record.source_blob == data
        assert record.source_file_size == len(data)


def test_file_upload_api_enforces_role_type_and_size(knowledge_repository, tmp_path):
    settings = Settings(
        run_store_path=tmp_path / "runs.json",
        knowledge_max_file_bytes=200_000,
    )
    client = TestClient(
        create_app(
            settings=settings,
            engine=make_engine(),
            knowledge_repository=knowledge_repository,
        )
    )
    form = {
        "document_key": "uploaded-policy",
        "title": "研发设备采购制度",
        "category_id": CATEGORY,
        "source_uri": "s3://policies/uploaded-policy.docx",
        "policy_version": "3",
        "effective_from": (utc_now() - timedelta(days=1)).isoformat(),
    }
    files = {"file": ("policy.docx", _docx_bytes(), DOCX_MEDIA_TYPE)}

    assert client.post(
        "/internal/knowledge/files", headers=EMPLOYEE_SESSION, data=form, files=files
    ).status_code == 403
    created = client.post(
        "/internal/knowledge/files", headers=ADMIN_SESSION, data=form, files=files
    )
    assert created.status_code == 200
    assert created.json()["version"]["source_file_name"] == "policy.docx"

    invalid_form = {**form, "document_key": "invalid-upload"}
    invalid = client.post(
        "/internal/knowledge/files",
        headers=ADMIN_SESSION,
        data=invalid_form,
        files={"file": ("policy.pdf", b"not a pdf", "application/pdf")},
    )
    assert invalid.status_code == 422
    assert invalid.json()["error"]["code"] == "FILE_TYPE_MISMATCH"

    oversized_form = {**form, "document_key": "oversized-upload"}
    oversized = client.post(
        "/internal/knowledge/files",
        headers=ADMIN_SESSION,
        data=oversized_form,
        files={"file": ("policy.txt", b"x" * 200_001, "text/plain")},
    )
    assert oversized.status_code == 422
    assert oversized.json()["error"]["code"] == "KNOWLEDGE_FILE_TOO_LARGE"


def test_openai_embedding_provider_batches_and_restores_response_order():
    calls = []

    class FakeEmbeddings:
        def create(self, **kwargs):
            calls.append(kwargs)
            data = []
            for index, text in reversed(list(enumerate(kwargs["input"]))):
                vector = [0.0] * KNOWLEDGE_EMBEDDING_DIMENSIONS
                vector[0] = float(text.removeprefix("item-"))
                data.append(SimpleNamespace(index=index, embedding=vector))
            return SimpleNamespace(data=data)

    provider = OpenAIEmbeddingProvider(
        api_key="fixture-key",
        batch_size=2,
        client=SimpleNamespace(embeddings=FakeEmbeddings()),
    )
    result = provider.embed(["item-1", "item-2", "item-3"])

    assert [vector[0] for vector in result] == [1.0, 2.0, 3.0]
    assert len(calls) == 2
    assert all(call["dimensions"] == KNOWLEDGE_EMBEDDING_DIMENSIONS for call in calls)
    assert provider.model == "openai:text-embedding-3-small:d256"


def test_embedding_backfill_is_tenant_scoped_batched_and_resumable(knowledge_repository):
    initial = PolicyKnowledgeService(knowledge_repository)
    ingested = initial.ingest(
        tenant_id="demo_cn",
        actor_user_id="knowledge_admin_001",
        roles=("KNOWLEDGE_STEWARD",),
        document_key="backfill-policy",
        title="研发设备采购制度",
        category_id=CATEGORY,
        source_uri="https://policies.example.test/backfill",
        policy_version="3",
        effective_from=utc_now() - timedelta(days=1),
        raw_content=(
            "第一条 内存标准\n研发笔记本不得低于32GB内存。\n"
            "第二条 保修标准\n保修不得少于36个月。"
        ),
    )
    initial.publish(
        ingested.document.document_id,
        tenant_id="demo_cn",
        roles=("KNOWLEDGE_STEWARD",),
    )
    replacement = PolicyKnowledgeService(
        knowledge_repository,
        embedding_provider=ReplacementEmbeddingProvider(),
    )

    preview = replacement.backfill_embeddings(tenant_id="demo_cn", dry_run=True)
    first = replacement.backfill_embeddings(
        tenant_id="demo_cn", batch_size=1, max_batches=1
    )
    finished = replacement.backfill_embeddings(tenant_id="demo_cn", batch_size=10)

    assert preview.pending_before >= 2
    assert preview.pending_after == preview.pending_before
    assert first.updated_chunks == 1
    assert first.pending_after == preview.pending_before - 1
    assert finished.pending_after == 0
    assert finished.updated_chunks == preview.pending_before - 1
    assert replacement.backfill_embeddings(tenant_id="other").updated_chunks == 0
    with knowledge_repository.session_factory() as session:
        count = session.execute(
            select(func.count(KnowledgeDocumentVersionRecord.version_id)).where(
                KnowledgeDocumentVersionRecord.tenant_id == "demo_cn"
            )
        ).scalar_one()
    assert count == 1
